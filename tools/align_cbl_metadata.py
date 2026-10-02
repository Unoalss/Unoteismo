"""Alinha a Falsa Folha de Rosto, a Folha de Rosto e a Ficha Catalográfica (CIP)
da Bíblia Sagrada (port-br) Unoteista (69 Livros - Sem Teofilo).docx
com o cadastro oficial da CBL:
- Título: Bíblia Sagrada Unoteísta
- Subtítulo: o caminho da verdade
- Titular / Autor: Arthur Santos
- Edição / Revisão: Caminho Da Verdade
- ISBN: 978-65-02-41194-0
- Paginação: 1730 p. ; 15,5 x 23,0 cm.
"""
import os
import re
import shutil
import zipfile
import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_PATH = os.path.join(ROOT, 'Bíblia Sagrada (port-br) Unoteista (69 Livros - Sem Teofilo).docx')

def run():
    doc = docx.Document(DOCX_PATH)

    # 1. Falsa Folha de Rosto (P0-P2)
    doc.paragraphs[0].text = 'BÍBLIA SAGRADA UNOTEÍSTA'
    doc.paragraphs[1].text = 'o caminho da verdade'
    doc.paragraphs[2].text = 'EDIÇÃO CAMINHO DA VERDADE'

    # 2. Folha de Rosto Oficial (P4-P8)
    doc.paragraphs[4].text = 'BÍBLIA SAGRADA UNOTEÍSTA'
    doc.paragraphs[5].text = 'o caminho da verdade'
    doc.paragraphs[6].text = 'Cânon Unoteísta de 69 Livros\nAntigo Testamento (42 livros) e Novo Testamento (27 livros)'
    doc.paragraphs[7].text = 'Arthur Santos\nEdição e Revisão: Caminho Da Verdade'
    doc.paragraphs[8].text = 'Brasil — 2026'

    # 3. Ficha Catalográfica (CIP) (P25-P29)
    # Localizar por '1730 p.'
    cip_idx = None
    for i, p in enumerate(doc.paragraphs[:50]):
        if '1730 p.' in p.text:
            cip_idx = i
            break

    if cip_idx:
        doc.paragraphs[cip_idx - 3].text = '        Bíblia Sagrada Unoteísta : o caminho da verdade / Arthur'
        doc.paragraphs[cip_idx - 2].text = '        Santos ; edição e revisão por Caminho Da Verdade. – 1. ed. –'
        doc.paragraphs[cip_idx - 1].text = '        Brasil : Edição Caminho Da Verdade, 2026.'
        doc.paragraphs[cip_idx].text     = '        1730 p. ; 15,5 x 23,0 cm.'
        doc.paragraphs[cip_idx + 1].text = '        ISBN 978-65-02-41194-0'
        print('Ficha Catalográfica alinhada perfeitamente com os dados CBL!')

    # 4. Assuntos e entradas secundárias
    for i, p in enumerate(doc.paragraphs[:50]):
        if 'I. Caminho Da Verdade.' in p.text:
            p.text = '        1. Bíblia – Versões em língua portuguesa. 2. Unoteísmo.\n        3. Teologia bíblica.\n        I. Santos, Arthur. II. Caminho Da Verdade. III. Título.'
            break

    doc.save(DOCX_PATH)
    print('Documento Word salvo com sucesso.')

    # 5. Garantir <Pages>1730</Pages> em docProps/app.xml
    tmp_path = DOCX_PATH + '.tmp_zip'
    with zipfile.ZipFile(DOCX_PATH, 'r') as zin:
        with zipfile.ZipFile(tmp_path, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == 'docProps/app.xml':
                    xml_str = data.decode('utf-8', errors='ignore')
                    if '<Pages>' in xml_str:
                        xml_str = re.sub(r'<Pages>\d+</Pages>', '<Pages>1730</Pages>', xml_str)
                    else:
                        xml_str = xml_str.replace('</Properties>', '<Pages>1730</Pages></Properties>')
                    data = xml_str.encode('utf-8')
                zout.writestr(item, data)
    shutil.move(tmp_path, DOCX_PATH)
    print('Metadados de páginas fixados em 1730.')

if __name__ == '__main__':
    run()
