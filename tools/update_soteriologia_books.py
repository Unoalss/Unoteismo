"""Atualiza os arquivos docx com a Soteriologia Unoteísta oficial:
1. Soteriologia_Unoteista.docx (documento dedicado)
2. Bíblia Sagrada (port-br) Unoteista (69 Livros - Sem Teofilo).docx (inserido na introdução antes de NOTA SOBRE ESTA EDIÇÃO)
"""
import os
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

SOTERIOLOGIA_DATA = [
    {
        "num": "1",
        "title": "1. A salvação procede de Deus",
        "paragraphs": [
            "Cremos que a salvação procede de Deus e é oferecida ao homem por meio de Cristo. O homem não produz em si mesmo a certeza necessária para a salvação, pois a fé é dada por Deus.",
            "A fé é a certeza concedida por Deus ao homem acerca daquilo que Ele promete. Entretanto, receber a fé não elimina a responsabilidade humana: o homem deve crer, isto é, responder conscientemente àquilo que Deus lhe deu."
        ],
        "chain": "Deus dá a fé → o homem recebe a certeza → o homem escolhe crer."
    },
    {
        "num": "2",
        "title": "2. Fé e crença",
        "paragraphs": [
            "Cremos que fé e crença não são a mesma coisa.",
            "A fé é a certeza que vem de Deus.",
            "A crença é a resposta consciente do homem diante dessa certeza.",
            "A salvação envolve, portanto:"
        ],
        "chain": "fé que vem de Deus + crença que procede do homem."
    },
    {
        "num": "3",
        "title": "3. A antiga aliança e a morte do Testador",
        "paragraphs": [
            "A salvação também está relacionada à validade da nova aliança.",
            "A Escritura declara que, onde há testamento, é necessário que intervenha a morte do testador:"
        ],
        "quote": ("“Porque onde há testamento, é necessário que intervenha a morte do testador.”", "— Hebreus 9:16"),
        "paragraphs_after": [
            "Por isso, Cristo é apresentado como Mediador de uma nova aliança, cuja validade está relacionada à sua morte e cuja finalidade é que os chamados recebam a promessa da herança eterna.",
            "A morte de Cristo, portanto, não deve ser entendida somente como morte física ou como remissão das transgressões, mas também como o acontecimento pelo qual a antiga ordem da aliança chega ao seu fim e a nova aliança é estabelecida."
        ]
    },
    {
        "num": "4",
        "title": "4. A morte do marido e a libertação da antiga aliança",
        "paragraphs": [
            "Paulo utiliza em Romanos 7:1–5 a figura do casamento para explicar a relação entre Israel, a Lei e a morte.",
            "Enquanto o marido vive, a mulher está ligada à lei do marido. Mas, ocorrendo a morte do marido, ela fica livre daquela lei e pode pertencer a outro.",
            "Dentro da compreensão Unoteísta, essa figura possui também uma dimensão relacionada à aliança entre Deus e Israel.",
            "Israel foi tomado como esposa de Deus, mas tornou-se infiel ao seu marido ao seguir outros deuses. A antiga aliança estabelecia esse vínculo e também tornava manifesta a transgressão da infidelidade.",
            "Como Deus é eterno, essa situação não poderia ser resolvida simplesmente pela morte definitiva de Deus.",
            "Por isso, a morte ocorreu por meio da encarnação: a Consciência divina, que é Deus, entrou na condição humana em Jesus Cristo e morreu segundo a carne.",
            "Assim, a morte de Cristo cumpre simultaneamente a figura do marido que morre e a figura do testador cuja morte torna válido o testamento.",
            "A morte ocorreu realmente, mas não significou a extinção de Deus, pois aquele que morreu segundo a carne foi ressuscitado."
        ],
        "chain": "O marido morreu → a antiga obrigação é encerrada → a nova aliança é estabelecida → Deus ressuscita → uma nova relação de aliança é oferecida."
    },
    {
        "num": "5",
        "title": "5. A nova aliança",
        "paragraphs": [
            "A ressurreição inaugura uma nova condição.",
            "Aquele que morreu ressuscitou e agora apresenta uma nova aliança fundamentada não na condenação da antiga aliança, mas na misericórdia e na benignidade de Deus.",
            "Por isso, a promessa anunciada pelo profeta encontra sua expressão na nova relação:"
        ],
        "quote": ("“E desposar-te-ei comigo para sempre; desposar-te-ei comigo em justiça, e em juízo, em benignidade e em misericórdias.”", "— Oséias 2:19"),
        "paragraphs_after": [
            "A nova aliança é, portanto, apresentada como uma nova união entre Deus e seu povo."
        ]
    },
    {
        "num": "6",
        "title": "6. Herdeiros e noiva em Cristo",
        "paragraphs": [
            "Na nova aliança, os que creem tornam-se herdeiros da promessa.",
            "A relação com Deus também é apresentada pela linguagem matrimonial: aqueles que pertencem a Cristo constituem a noiva.",
            "Assim, em Cristo, os homens são simultaneamente:"
        ],
        "chain": "filhos de Deus → herdeiros da promessa → membros da nova aliança → noiva de Deus.",
        "paragraphs_after": [
            "A herança não é conquistada por mérito humano, mas recebida mediante a promessa de Deus."
        ]
    },
    {
        "num": "7",
        "title": "7. O mérito humano na salvação",
        "paragraphs": [
            "O homem não possui mérito pela origem da salvação, porque a salvação, a fé e a graça procedem de Deus.",
            "Se todo o processo da salvação fosse produzido pelo homem, então o homem poderia atribuir a si mesmo o mérito de sua salvação.",
            "Mas a salvação não é 100% obra humana.",
            "Deus oferece a salvação, concede a fé e, depois do selamento, sustenta o homem nela.",
            "O único processo relacionado à perda da salvação que é 100% atribuído ao homem é a rejeição da salvação que Deus lhe ofereceu."
        ],
        "chain": "O homem não produz a salvação; o homem pode rejeitá-la."
    },
    {
        "num": "8",
        "title": "8. A possibilidade de rejeição",
        "paragraphs": [
            "Antes do selamento pelo Espírito da Verdade, o homem ainda pode rejeitar a graça que recebeu.",
            "Aquele que recebeu a fé pode escolher não crer, abandonar a fé ou rejeitar aquilo que Deus lhe ofereceu.",
            "Nesse estágio, existe a possibilidade de cair da graça e perder a salvação."
        ]
    },
    {
        "num": "9",
        "title": "9. O selamento pelo Espírito da Verdade",
        "paragraphs": [
            "Quando ocorre a relação entre fé e crença, Deus sela o homem com o Espírito da Verdade, que é o penhor da salvação.",
            "A partir desse momento, Deus passa a sustentar o homem na fé."
        ],
        "subsections": [
            ("Antes do selamento:", "Deus oferece → Deus concede a fé → o homem crê ou rejeita."),
            ("Depois do selamento:", "Deus sustenta → Deus opera → Deus guia → Deus conduz até o fim.")
        ],
        "paragraphs_after": [
            "O Espírito da Verdade ensina, guia, corrige e conduz o homem na verdade."
        ]
    },
    {
        "num": "10",
        "title": "10. Deus efetua o querer e o realizar",
        "paragraphs": [
            "Depois do selamento, a perseverança não depende exclusivamente da capacidade humana.",
            "Deus passa a efetuar no homem o querer e o realizar, conduzindo-o conforme a sua vontade.",
            "Por isso, aquele que foi selado não perde novamente a salvação, porque Deus mesmo passa a sustentar nele a fé e conduzi-lo até a consumação."
        ]
    },
    {
        "num": "11",
        "title": "11. A filiação",
        "paragraphs": [
            "Quando o homem recebe a fé, crê e é selado pelo Espírito da Verdade, torna-se filho de Deus.",
            "Essa filiação já é uma realidade espiritual, mas sua manifestação plena ainda está no futuro.",
            "Os filhos aguardam a manifestação daquilo que serão na ressurreição."
        ]
    },
    {
        "num": "12",
        "title": "12. A consumação da salvação",
        "paragraphs": [
            "Na ressurreição, os filhos de Deus receberão corpos espirituais incorruptíveis.",
            "Assim como Cristo foi revestido de um corpo glorioso, os seus serão transformados à semelhança do seu corpo glorioso.",
            "Então a filiação que hoje existe pela fé será plenamente manifestada.",
            "Os filhos verão como Ele é e serão semelhantes a Ele."
        ]
    },
    {
        "num": "13",
        "title": "13. O Pai conhecido por meio do Cordeiro",
        "paragraphs": [
            "O Deus invisível será conhecido por seus filhos por meio daquele que é sua manifestação gloriosa: o Cordeiro.",
            "Cristo é o meio pelo qual os filhos contemplam o Pai.",
            "Assim, a esperança final da salvação é entrar plenamente na nova aliança, receber a herança prometida, participar da união com Deus como sua noiva e contemplar o Pai por meio do Cordeiro."
        ]
    }
]

RESUMO_PONTOS = [
    "Deus oferece a salvação.",
    "Deus concede a fé.",
    "O homem escolhe crer.",
    "A fé + a crença estabelecem a relação de salvação.",
    "Cristo morre, cumprindo a necessidade da morte do Testador e a figura do marido que morre.",
    "A antiga aliança chega ao seu fim e a nova aliança é estabelecida.",
    "Cristo ressuscita.",
    "O homem é selado pelo Espírito da Verdade.",
    "Deus passa a efetuar nele o querer e o realizar.",
    "A salvação não é mais perdida.",
    "O homem torna-se filho e herdeiro.",
    "Na ressurreição, recebe um corpo glorioso.",
    "Torna-se semelhante a Cristo e contempla o Pai por meio do Cordeiro.",
    "Assim, a nova aliança é uma aliança de justiça, benignidade e misericórdia."
]


def add_soteriologia_elements(doc_or_insert_func):
    """Gera a lista de elementos formatados da Soteriologia."""
    elements = []
    
    # Título Principal
    elements.append(('TITLE', 'SOTERIOLOGIA UNOTEÍSTA'))
    elements.append(('SUBTITLE', 'A Doutrina Bíblica da Salvação: Da Graça e Concessão da Fé ao Selamento e Consumação Eterna'))
    elements.append(('SPACE', ''))
    
    for sec in SOTERIOLOGIA_DATA:
        elements.append(('SEC_TITLE', sec['title']))
        for p in sec.get('paragraphs', []):
            elements.append(('PARA', p))
        if 'quote' in sec:
            elements.append(('QUOTE', sec['quote']))
        if 'chain' in sec:
            elements.append(('CHAIN', sec['chain']))
        if 'subsections' in sec:
            for sub_lbl, sub_txt in sec['subsections']:
                elements.append(('SUB_BLOCK', (sub_lbl, sub_txt)))
        for p in sec.get('paragraphs_after', []):
            elements.append(('PARA', p))
        elements.append(('SPACE_SEC', ''))
        
    elements.append(('RESUMO_TITLE', 'Resumo da Soteriologia Unoteísta'))
    for pt in RESUMO_PONTOS:
        elements.append(('RESUMO_ITEM', pt))
    elements.append(('SPACE', ''))
    
    return elements


def format_paragraph(p, elem_type, data):
    p.paragraph_format.line_spacing = 1.15
    if elem_type == 'TITLE':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(data)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.bold = True
    elif elem_type == 'SUBTITLE':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(data)
        run.font.name = 'Bookman Old Style'
        run.font.size = Pt(11)
        run.italic = True
    elif elem_type == 'SEC_TITLE':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(data)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.bold = True
    elif elem_type == 'PARA':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(data)
        run.font.name = 'Bookman Old Style'
        run.font.size = Pt(10.5)
    elif elem_type == 'QUOTE':
        quote_text, ref_text = data
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.right_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        r1 = p.add_run(quote_text + '\n')
        r1.font.name = 'Bookman Old Style'
        r1.font.size = Pt(10)
        r1.bold = True
        r2 = p.add_run(ref_text)
        r2.font.name = 'Bookman Old Style'
        r2.font.size = Pt(9.5)
        r2.italic = True
    elif elem_type == 'CHAIN':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(data)
        run.font.name = 'Bookman Old Style'
        run.font.size = Pt(10.5)
        run.bold = True
    elif elem_type == 'SUB_BLOCK':
        sub_lbl, sub_txt = data
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(sub_lbl + ' ')
        r1.font.name = 'Bookman Old Style'
        r1.font.size = Pt(10)
        r1.bold = True
        r2 = p.add_run(sub_txt)
        r2.font.name = 'Bookman Old Style'
        r2.font.size = Pt(10)
        r2.italic = True
    elif elem_type == 'RESUMO_TITLE':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(data)
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.bold = True
    elif elem_type == 'RESUMO_ITEM':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_bullet = p.add_run('• ')
        r_bullet.bold = True
        r_bullet.font.name = 'Arial'
        run = p.add_run(data)
        run.font.name = 'Bookman Old Style'
        run.font.size = Pt(10.5)
        run.bold = True
    elif elem_type in ('SPACE', 'SPACE_SEC'):
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)


def update_standalone_soteriologia():
    target_path = os.path.join(ROOT, 'Soteriologia_Unoteista.docx')
    bak_path = target_path + '.bak'
    if os.path.exists(target_path) and not os.path.exists(bak_path):
        shutil.copy2(target_path, bak_path)
    
    doc = docx.Document()
    # Margens padrão 2.5cm
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.98)
        s.bottom_margin = Inches(0.98)
        s.left_margin = Inches(0.98)
        s.right_margin = Inches(0.98)
        
    elements = add_soteriologia_elements(doc)
    for elem_type, data in elements:
        p = doc.add_paragraph()
        format_paragraph(p, elem_type, data)
        
    doc.save(target_path)
    print(f'Atualizado: {target_path}')


def update_bible_book():
    bible_path = os.path.join(ROOT, 'Bíblia Sagrada (port-br) Unoteista (69 Livros - Sem Teofilo).docx')
    if not os.path.exists(bible_path):
        print(f'Arquivo não encontrado: {bible_path}')
        return
        
    bak_path = bible_path + '.soteriologia.bak'
    if not os.path.exists(bak_path):
        shutil.copy2(bible_path, bak_path)
        
    doc = docx.Document(bible_path)
    
    # Localizar o ponto de inserção: logo após o Artigo 12 da Confissão de Fé (antes de NOTA SOBRE ESTA EDIÇÃO)
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if 'Artigo 12' in p.text:
            # O parágrafo seguinte é o texto do Artigo 12
            target_idx = i + 2  # insere logo após o texto do artigo 12
            break
            
    if target_idx is None:
        print('AVISO: Artigo 12 não localizado na Bíblia.')
        return
        
    print(f'Localizado Artigo 12. Ponto de inserção: parágrafo {target_idx}')
    
    # Inserir elementos da Soteriologia antes de NOTA SOBRE ESTA EDIÇÃO
    # Em python-docx, p.insert_paragraph_before() insere antes de um parágrafo existente
    ref_p = doc.paragraphs[target_idx]
    
    # Vamos conferir o texto do parágrafo de referência
    print(f'Inserindo antes de: {ref_p.text[:60]!r}')
    
    elements = add_soteriologia_elements(doc)
    for elem_type, data in elements:
        new_p = ref_p.insert_paragraph_before()
        format_paragraph(new_p, elem_type, data)
        
    doc.save(bible_path)
    print(f'Atualizado com sucesso: {bible_path}')


if __name__ == '__main__':
    update_standalone_soteriologia()
    update_bible_book()
